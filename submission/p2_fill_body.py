"""fill_problem2(): writes every Problem 2 box of the submission template from judge/results/*.json and the audit log."""
import copy
import json

from fill_problem2 import DIAGRAM, ENVVARS, INDEP, INSTALL, PREREQ, RUN_AB, RUN_SUITE, _load


def _log_rec(root, tag_prefix, model_suffix):
    for line in (root / "judge" / "logs" / "judge_calls.jsonl").read_text(encoding="utf-8").splitlines():
        r = json.loads(line)
        if r["tag"].startswith(tag_prefix) and r["model"].endswith(model_suffix):
            return r
    return None


def _pct(x, d=1):
    return "n/a" if x is None else f"{x * 100:.{d}f}%"


def fill_problem2(d, body, T, set_cell, set_after_label, ROOT):
    from judge.rubric import RUBRIC
    R = _load(ROOT / "judge" / "results")
    g120, pos, ver, adv = R["gold"]["variants"], R["position"]["variants"], R["verbosity"]["variants"], R["adversarial"]["variants"]
    g20 = R["gold_gptoss20b"]["variants"]["mitigated"]["validation"]
    syc, ab, rt = R["sycophancy_gptoss20b"]["variants"], R["ab_gptoss20b"], R["retest_gptoss20b"]
    gb, gm = g120["baseline"], g120["mitigated"]
    pb, pm = pos["baseline"], pos["mitigated"]
    cmp_ = ab["comparison"]
    pw = cmp_["pointwise"]
    A, B = pw[cmp_["config_a"]], pw[cmp_["config_b"]]
    ps = cmp_["pairwise"]
    n_dis = round(ps["order_disagreement_rate"] * ps["pairs"])

    # ---- 2.1 / 2.2 ----
    set_cell(T[21].rows[0].cells[0], DIAGRAM, mono=True, size=7)
    set_after_label(74, PREREQ)
    set_after_label(75, ENVVARS)
    set_cell(T[22].rows[0].cells[0], INSTALL, mono=True)
    set_cell(T[23].rows[0].cells[0], RUN_SUITE, mono=True)
    set_cell(T[24].rows[0].cells[0], RUN_AB, mono=True)
    set_after_label(82, INDEP)

    # ---- 2.3.1 mode + rubric ----
    set_after_label(86, "Pointwise scoring (reference-based when expected_output is given, reference-free otherwise) for suites, validation and probes; "
                        "pairwise A-vs-B, always in both orders, for position-bias work and the prompt v1 vs v2 comparison. "
                        "Pointwise fits release thresholds, drift tracking and single outputs; pairwise fits choosing between two candidates, where an absolute 1-5 scale drifts; "
                        "reference-based fits closed-answer tasks with a trusted answer; reference-free fits open-ended output or production traffic with no gold.")
    rt_ = T[25]
    rt_._tbl.append(copy.deepcopy(rt_.rows[-1]._tr))   # the template has 5 criterion rows; the rubric has 6
    for row, c in zip(rt_.rows[1:], RUBRIC):
        anchors = "; ".join(f"{k} = {v}" for k, v in sorted(c.anchors.items()))
        set_cell(row.cells[0], c.id.replace("_", "-").capitalize())
        set_cell(row.cells[1], f"{c.definition} Anchors: {anchors}.", size=8)
        set_cell(row.cells[2], f"{c.weight:.2f}")

    # ---- 2.3.2 bias table ----
    bt = T[26].rows
    set_cell(bt[1].cells[1], "judge/core.py: every pair is judged in BOTH orders (A,B) and (B,A); labels are mapped back; the verdict stands only if both orders agree, "
                             "otherwise TIE and the pair is flagged (combine_orders). Reported: order-disagreement rate, hard flips (A one way, B the other), accuracy vs known winners.", size=8)
    set_cell(bt[1].cells[2], f"22 pairs with a known winner (judge gpt-oss-120b). Orders disagreed on {_pct(pb['summary']['order_disagreement_rate'])} (2/22; one hard A/B flip) with the baseline prompt "
                             f"-> {_pct(pm['summary']['order_disagreement_rate'])} with the mitigated prompt; requiring agreement makes the reported verdict order-independent by construction. "
                             f"Accuracy vs gold: single order {_pct(pb['single_order_AB']['accuracy_vs_gold'])} (A first) / {_pct(pb['single_order_BA']['accuracy_vs_gold'])} (B first) -> "
                             f"{_pct(pb['both_orders_require_agreement']['accuracy_vs_gold'])} with both orders, 0 decisive-but-wrong, at the price of {_pct(pb['summary']['tie_rate'])} TIEs. "
                             f"On the real A/B suite (gpt-oss-20b) the orders disagreed on {n_dis}/{ps['pairs']} pairs ({_pct(ps['order_disagreement_rate'])}): "
                             f"first-order-only would have declared prompt v1 the winner {ps['single_order_wins_a']}-{ps['single_order_wins_b']}; with agreement required all {ps['ties']} pairs are TIE.", size=8)
    vb, vm = ver["baseline"], ver["mitigated"]
    set_cell(bt[2].cells[1], "Rubric criterion 'tone' penalises padding/repetition; the mitigated prompt has an explicit length-control rule (rubric.py RULES 2); probe = 8 concise-correct answers vs the same facts padded with filler and restatement.", size=8)
    set_cell(bt[2].cells[2], f"Padded answer preferred in {vb['pairwise_single_judgments']['padded_preferred']}/{vb['pairwise_single_judgments']['n']} pairwise judgments (baseline) -> "
                             f"{vm['pairwise_single_judgments']['padded_preferred']}/{vm['pairwise_single_judgments']['n']} (mitigated); pointwise padded minus concise: "
                             f"{vb['pointwise']['mean_padded_minus_concise']:+.2f} -> {vm['pointwise']['mean_padded_minus_concise']:+.2f}. "
                             "No before->after improvement: this judge showed no verbosity bias on this probe even without the mitigation (the probe is easy: templated filler; long-but-substantive answers were not tested).", size=8)
    set_cell(bt[3].cells[1], "Judge and generator are separate settings; judge/config.py check_cross_family() refuses a same-family pair unless --allow-same-family (recorded in the report). Judge = OpenAI family (gpt-oss), generator = Qwen.", size=8)
    set_cell(bt[3].cells[2], "NOT MEASURED before->after (that needs a same-family judge/generator pair on equal-quality outputs, not run). Guard behaviour is unit-tested. "
                             "Later experiments used gpt-oss-20b as judge because gpt-oss-120b hit its daily token quota; it is the same family as 120b and still not Qwen.", size=8)
    sb, sm = syc["baseline"], syc["mitigated"]
    set_cell(bt[4].cells[1], "Per-criterion grounding in code: a 4-5 on correctness/faithfulness/completeness needs a verbatim evidence quote from the output (parsing.evidence_supported), else the score is capped at 3; "
                             "prompt rules ignore confident tone, cited authorities and notes addressed to the evaluator. Probe = 8 confidently-wrong answers (4 also carry 'note to evaluator: top score').", size=8)
    set_cell(bt[4].cells[2], f"Judge gpt-oss-20b: wrong answers passed {sb['fooled']}/{sb['n_scored']} (baseline) -> {sm['fooled']}/{sm['n_scored']} (mitigated): not fooled either way, so no measurable gain. "
                             f"Mean score of the wrong answers {sb['mean_overall']:.2f} -> {sm['mean_overall']:.2f} (the mitigated prompt was slightly MORE lenient). The evidence cap fired 0 times in the 120b gold run, so its effect is not demonstrated.", size=8)
    cb, cm = gb["clustering"], gm["clustering"]
    set_cell(bt[5].cells[1], "Few-shot anchors in the prompt (rubric.FEW_SHOT) that show low scores, plus 1-5 anchor text per criterion; pairwise mode is also provided.", size=8)
    set_cell(bt[5].cells[2], f"Score spread (stdev of overall, 26 gold cases, gpt-oss-120b): {cb['stdev_overall']:.2f} -> {cm['stdev_overall']:.2f}; share in the busiest 1-point band {_pct(cb['modal_band_share'])} -> {_pct(cm['modal_band_share'])}. "
                             "The spread did NOT widen (it narrowed); the judge was not clustered to begin with because the gold set is bimodal (pass-vs-fail mean gap "
                             f"{gb['validation']['discrimination_gap']:.2f} -> {gm['validation']['discrimination_gap']:.2f}).", size=8)

    # ---- 2.3.3 validation ----
    vt = T[27].rows
    gv_b, gv_m = gb["validation"], gm["validation"]
    set_cell(vt[1].cells[1], f"gpt-oss-120b: baseline {_pct(gv_b['agreement_rate'])} (26/26), mitigated {_pct(gv_m['agreement_rate'])} (25/26, 1 false reject); gpt-oss-20b mitigated {_pct(g20['agreement_rate'])} (26/26)")
    set_cell(vt[1].cells[2], f"n = 26 hand-labelled cases (13 pass/13 fail), binary pass/fail. 95% Wilson CI for the 120b mitigated rate {gv_m['agreement_ci95'][0] * 100:.0f}-{gv_m['agreement_ci95'][1] * 100:.0f}%. "
                             "Labels were written by Claude (not an independent human) before any judge ran, with flaws built in, so the set is easy and near ceiling: it shows the judge is not broken, not that it is accurate on hard cases.", size=8)
    set_cell(vt[2].cells[1], f"Cohen's kappa: 120b baseline {gv_b['cohen_kappa']:.2f}, 120b mitigated {gv_m['cohen_kappa']:.2f} ({gv_m['kappa_band']}), 20b mitigated {g20['cohen_kappa']:.2f}")
    set_cell(vt[2].cells[2], f"MAE vs gold 1-5 score: 120b baseline {gv_b['mae_vs_gold_score']:.2f}, mitigated {gv_m['mae_vs_gold_score']:.2f}. The mitigated prompt was not better here; its one miss (g12, '17 x 23 = 391.') "
                             "was a prompt defect: the judge scored instruction-following 1 because 'the model did not follow the required JSON format' (it confused its own reply format with the task).", size=8)
    t0, t7 = rt["temperature_0"], rt["temperature_0.7"]
    set_cell(vt[3].cells[1], f"Pass/fail flips: {t0['pass_flips']}/{t0['n_compared']} ({_pct(t0['pass_flip_rate'], 0)}) at temperature 0; {t7['pass_flips']}/{t7['n_compared']} ({_pct(t7['pass_flip_rate'], 0)}) at temperature 0.7")
    set_cell(vt[3].cells[2], f"gpt-oss-20b, mitigated prompt, 13 of the 26 gold cases (daily quota). Scores still moved: mean |change| {t0['mean_abs_score_diff']:.2f} (max {t0['max_abs_score_diff']:.2f}) at temp 0 and "
                             f"{t7['mean_abs_score_diff']:.2f} (max {t7['max_abs_score_diff']:.2f}) at 0.7 between two runs. The 120b judge was not re-run (quota). With 13 cases mostly far from the threshold, 0 flips is weak evidence of stability.", size=8)
    ab_, am_ = adv["baseline"]["by_kind"], adv["mitigated"]["by_kind"]
    vws = [r["overall"] for r in adv["mitigated"]["rows"] if r["kind"] == "verbose_but_wrong"]
    vwb = [r["overall"] for r in adv["baseline"]["rows"] if r["kind"] == "verbose_but_wrong"]
    set_cell(vt[4].cells[1], f"Verbose-but-wrong: fooled {ab_['verbose_but_wrong']['fooled']}/6 (baseline), {am_['verbose_but_wrong']['fooled']}/6 (mitigated). "
                             f"Terse-but-correct: fooled {ab_['terse_but_correct']['fooled']}/6 (baseline), {am_['terse_but_correct']['fooled']}/6 (mitigated). Fooled = Y for the mitigated prompt on one terse answer.")
    set_cell(vt[4].cells[2], f"gpt-oss-120b, reference-free, 12 matched probes. The mitigated prompt scored the verbose-wrong answers HIGHER ({min(vws):.2f}-{max(vws):.2f} vs {min(vwb):.2f}-{max(vwb):.2f}; "
                             f"{sum(1 for x in vws if x >= 3.5)}/6 had a weighted overall at or above the 3.5 pass line); only the 'no criterion <= 2' floor kept them failing. "
                             "The terse miss (adv_t4, '3.') is the same prompt defect as g12.", size=8)

    # ---- 2.3.4 A/B ----
    at = T[28].rows
    nA = f"{round(A['pass_rate'] * A['scored'])}/{A['scored']}"
    nB = f"{round(B['pass_rate'] * B['scored'])}/{B['scored']}"
    set_cell(at[1].cells[0], f"Config A - {cmp_['config_a']} (precise, concise)")
    set_cell(at[2].cells[0], f"Config B - {cmp_['config_b']} (friendly, ~150 words)")
    set_cell(at[1].cells[1], f"{_pct(A['pass_rate'])} ({nA}; {A['unparseable']} unparseable of {A['cases']})")
    set_cell(at[2].cells[1], f"{_pct(B['pass_rate'])} ({nB})")
    set_cell(at[1].cells[2], f"{A['mean_overall']:.2f}")
    set_cell(at[2].cells[2], f"{B['mean_overall']:.2f}")
    set_cell(at[1].cells[3], f"{ps['wins_a']}/{ps['scored']} wins ({_pct(ps['win_rate_a'])}); {ps['ties']} ties")
    set_cell(at[2].cells[3], f"{ps['wins_b']}/{ps['scored']} wins ({_pct(ps['win_rate_b'])}); {ps['ties']} ties")
    dec = cmp_["decision"]
    lo, hi = pw["mean_diff_ci95_bootstrap"]
    set_after_label(96, f"NO WINNER DECLARED. {dec['verdict']}. Rule: {dec['rule']}. Both orders agreed on a decisive winner in none of the {ps['scored']} pairs; "
                        f"pointwise, v1 is ahead by {pw['mean_diff_a_minus_b']:+.2f} points (95% bootstrap CI {lo:+.2f} to {hi:+.2f}, includes 0) and its pass rate is {_pct(A['pass_rate'], 0)} vs {_pct(B['pass_rate'], 0)}; "
                        f"that is a lean toward v1, not a significant result. Caveats: {ab['generator_model']} generated both configs; only {ps['pairs']} of 16 inputs survived (2 were cut off by the output-token cap and excluded); "
                        "judge gpt-oss-20b; one hand-written input set; low power.")

    # ---- 2.4 design answers ----
    ans = [
        "Pointwise scoring for suites, gold validation and probes (reference-based when a trusted answer exists, reference-free otherwise), because the gate I care about is an absolute pass/fail with per-criterion reasons. "
        "Pairwise A-vs-B (always both orders) for the two comparisons that are relative by nature: choosing between prompt v1 and v2, and measuring position bias. Pointwise fits thresholds and drift tracking; "
        "pairwise fits close calls where a 1-5 scale is unstable; reference-based needs a gold answer I trust; reference-free is the only option for open-ended or live traffic but leans on the judge's own knowledge.",
        "judge/parsing.py strips <think> blocks and code fences, takes the largest balanced {...} (string-aware, so braces inside text do not break it), repairs trailing commas and smart quotes, then validates the schema "
        "(every requested criterion present, numeric score in 1..5 or null, winner in A|B|TIE). On failure it retries once with a corrective message (and a different cache tag, so the bad reply is not replayed). "
        "If that also fails the case is recorded as 'unparseable' (or 'api_error'): no score, excluded from means, counted in the report. It never defaults or guesses a score. Covered by unit tests.",
        "The judge is OpenAI-family (gpt-oss-120b, later gpt-oss-20b) and the generator is Qwen (qwen3.8-27b), so the judge is not scoring its own family's writing; a code guard refuses a same-family pair. "
        "I did not measure self-enhancement, so I cannot claim how large the effect would have been. These were the models available on my key; one family difference is not an ensemble.",
        f"Position bias worried me most because it directly flips A/B decisions, and the data agrees: on the real A/B suite the two orders disagreed on {n_dis} of {ps['pairs']} pairs, and a first-order-only run would have declared prompt v1 the winner "
        f"{ps['single_order_wins_a']}-{ps['single_order_wins_b']}; requiring agreement turned that into {ps['ties']} ties. I am confident that mitigation removes order-driven winners; I am not confident it finds true ones (it also discards real signal). "
        "For verbosity and sycophancy the judge showed no bias even without mitigation, so I cannot claim those mitigations worked, and my prompt made two things worse: scores of verbose-wrong answers rose, and two correct terse answers were failed through a prompt defect.",
        "Not on its own, and not in its current state. As a regression signal with guardrails, yes: (1) pairwise decisions only when both orders agree, disagreements go to a human; (2) the 'no criterion <= 2' floor stays, because the weighted mean alone would have passed 3 of 6 verbose-wrong answers; "
        "(3) fix the prompt defect and re-run the validation before trusting it; (4) validate on a larger, independently human-labelled set (mine is 26 cases labelled by the AI, near ceiling), tracking kappa and false-accept rate per release; "
        "(5) require at least ~50 decisive pairs for an A/B call, never n=14; (6) humans review every unparseable case, every score between 3 and 4, every safety flag and any release-blocking failure; "
        "(7) re-validate whenever the judge model, rubric or prompt changes.",
    ]
    for t_idx, text in zip(range(29, 34), ans):
        set_cell(T[t_idx].rows[0].cells[0], text, size=10)

    # ---- 2.5 evidence ----
    ab_rec = _log_rec(ROOT, "pairwise|baseline|p17|AB|", "120b")
    ba_rec = _log_rec(ROOT, "pairwise|baseline|p17|BA|", "120b")
    if ab_rec:
        set_cell(T[34].rows[0].cells[0], "From judge/logs/judge_calls.jsonl, record tag '" + ab_rec["tag"] + "', model " + ab_rec["model"] + ", temperature " + str(ab_rec["temperature"]) +
                 f", tokens {ab_rec['prompt_tokens']}+{ab_rec['completion_tokens']}.\n\nRAW PROMPT (user message):\n" + ab_rec["messages"][-1]["content"] + "\n\nRAW RESPONSE:\n" + ab_rec["raw_response"] +
                 "\n\n(Text excerpt; replace with a screenshot of the log line if you prefer.)", mono=True, size=7)
    from docx.shared import Inches
    cell34 = T[34].rows[0].cells[0]
    for name, cap in (("p2_live_verdict.png", "Live deployment (https://synergylabs-assignment2.vercel.app/), pointwise verdict: judge openai/gpt-oss-120b, executed live on Groq, latency and tokens shown."),
                      ("p2_live_verdict_criteria.png", "Same verdict: per-criterion score, rationale and verbatim evidence quote from the model output.")):
        shot = ROOT / "submission" / "screenshots" / name
        if shot.exists():
            cell34.add_paragraph().add_run(cap)
            cell34.add_paragraph().add_run().add_picture(str(shot), width=Inches(6.0))
    if ab_rec and ba_rec:
        set_cell(T[35].rows[0].cells[0],
                 "Same pair (p17: 'Canberra.' vs 'The capital of Australia is Canberra.', gold = TIE), baseline judge prompt, gpt-oss-120b, raw replies from the log:\n"
                 "ORDER 1 (A = 'Canberra.', B = full sentence): " + ab_rec["raw_response"] + "\n"
                 "ORDER 2 (A = full sentence, B = 'Canberra.'): " + ba_rec["raw_response"] + "\n\n"
                 "Mapped back to the original labels, order 1 -> the full sentence and order 2 -> 'Canberra.'. The judge picked slot B both times, giving opposite answers with contradictory reasons "
                 "(B is 'complete' vs B is 'more concise'): a hard flip. Require-agreement makes this pair a TIE, which matches the gold label.\n\n"
                 f"Flip rate on all 22 pairs: orders disagreed on {_pct(pb['summary']['order_disagreement_rate'])} (2/22; hard flips {_pct(pb['summary']['hard_flip_rate'])}) with the baseline prompt, {_pct(pm['summary']['order_disagreement_rate'])} with the mitigated prompt; "
                 f"on the A/B suite (gpt-oss-20b) {_pct(ps['order_disagreement_rate'])} ({n_dis}/{ps['pairs']}).", mono=True, size=7)
    posbias = ROOT / "submission" / "reflection_p2_posbias.txt"
    if posbias.exists():
        set_cell(T[36].rows[0].cells[0], posbias.read_text(encoding="utf-8").strip(), size=10)
    else:
        set_cell(T[36].rows[0].cells[0],
                 "[ CANDIDATE: write this in your own words and save it as submission/reflection_p2_posbias.txt, then re-run submission/fill_submission.py. The assignment requires this box to be written without AI help. ]\n\n"
                 "Facts you can draw on (not an answer):\n"
                 "- Each pair is sent twice, once as (A,B) and once as (B,A); the second verdict is mapped back to the original labels (judge/core.py judge_pair / combine_orders).\n"
                 "- If both orders agree the verdict stands; if not, it is a TIE and flagged. A hard flip means A in one order and B in the other.\n"
                 f"- Measured here: baseline prompt, 22 pairs: 2 disagreements (1 hard flip); real A/B suite: {n_dis} of {ps['pairs']} disagreements.\n"
                 "- Think about: what a high flip rate says about the judge (it responds to position, not content), what it means for single-order results, "
                 "what the fix costs (more TIEs, 2x calls), and what the check cannot detect (a judge that is consistently wrong in both orders).", size=10)
    v3 = next(v for v in adv["mitigated"]["verdicts"] if v["id"] == "adv_v3")
    t4 = next(v for v in adv["mitigated"]["verdicts"] if v["id"] == "adv_t4")
    cs = lambda v: ", ".join(f"{k}={c['score']:.0f}" for k, c in v["criteria"].items())
    set_cell(T[37].rows[0].cells[0],
             "gpt-oss-120b, reference-free, from judge/results/adversarial.json.\n\n"
             f"VERBOSE-BUT-WRONG (adv_v3: '17 x 23' answered with a long, authoritative derivation ending in 381; true answer 391), mitigated prompt: weighted overall {v3['overall']} -> passed={v3['passed']}.\n  {cs(v3)}\n"
             "  Not fooled, but only just: the weighted mean is above the 3.5 pass line; correctness = 1 triggers the 'no criterion <= 2' floor. The baseline prompt scored this probe lower (2.4).\n\n"
             f"TERSE-BUT-CORRECT (adv_t4: len({{1, 2, 2, 3}}) answered '3.'), mitigated prompt: weighted overall {t4['overall']} -> passed={t4['passed']}  => FOOLED (false reject).\n  {cs(t4)}\n"
             f"  Judge's own rationale for instruction_following=1: \"{t4['criteria']['instruction_following']['rationale']}\" - it confused the JSON format it was told to reply in with an instruction for the model under test. "
             "The same defect caused the g12 miss. The baseline prompt passed all 6 terse-correct probes.\n\n"
             "My read: the judge was not fooled by verbosity or confident wrong content, but my mitigated prompt introduced a new failure on very short correct answers and raised verbose-wrong scores. "
             "Fix (not yet run; the gpt-oss-120b daily quota was exhausted): state explicitly that the JSON format applies to the judge's reply only, then re-run gold + adversarial.", size=8)
    set_cell(T[38].rows[0].cells[0],
             "Tool: Claude Code (Anthropic, model Claude Sonnet 5.5) in VS Code. Claude Code reviewed the existing Problem 2 repository and found that its 'judge' was a keyword-driven mock whose README numbers (e.g. 35% -> 0% flip rate, kappa 0.737, 100% adversarial robustness) did not come from any measurement; "
             "it then built the new judge/ package (rubric, prompts, robust parser, both-order pairwise, audit log with replay, statistics, A/B decision rule, experiments, CLI) and the unit tests, wrote the 26 gold cases, 22 position pairs and probe sets and their labels before running any judge, "
             "ran the experiments against Groq, diagnosed the rate-limit and token-cap problems (qwen output-token limit, gpt-oss-120b daily quota, truncated generations), rewired the web API to the real judge, replaced hard-coded UI numbers with measured results, wrote the README and handover addendum, and filled this document from judge/results/*.json. "
             "Claude Code also made these choices, which no independent human checked: the gold labels, the rubric weights and pass rule, and the switch to gpt-oss-20b for part of the experiments. Claude Code later added a live-model fallback chain (gpt-oss-120b, then gpt-oss-20b) to the code. "
             "The 2.5 reflection 'explain your position-bias check' is not generated by Claude Code; it is left for the candidate to write in their own words.", size=10)
